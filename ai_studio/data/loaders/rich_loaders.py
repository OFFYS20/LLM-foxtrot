"""PDF, DOCX and EPUB loaders (each needs an optional dependency)."""

from __future__ import annotations

import re
from pathlib import Path

from ai_studio.core.errors import DependencyMissingError, IngestionError
from ai_studio.data.loaders.base import DocumentLoader, LoadedDocument


class PDFLoader(DocumentLoader):
    extensions = {".pdf"}
    doc_type = "pdf"

    def load(self, path: Path, *, max_rows: int | None = None) -> LoadedDocument:
        try:
            import pymupdf as fitz
        except ImportError:
            try:
                import fitz  # older PyMuPDF releases
            except ImportError as exc:
                raise DependencyMissingError("pymupdf", "PDF import", install="pymupdf") from exc

        try:
            document = fitz.open(path)
        except Exception as exc:  # noqa: BLE001 - corrupt/encrypted files
            raise IngestionError(
                f"Could not open {path.name}: {exc}",
                hint="The file may be corrupt or password-protected.",
            ) from exc

        pages: list[str] = []
        empty_pages = 0
        try:
            if document.needs_pass:
                raise IngestionError(
                    f"{path.name} is password-protected.",
                    hint="Remove the password and import it again.",
                )
            for page in document:
                try:
                    text = page.get_text("text") or ""
                except Exception:  # noqa: BLE001 - a bad page must not kill the import
                    text = ""
                if not text.strip():
                    empty_pages += 1
                pages.append(text)

            metadata = document.metadata or {}
            title = (metadata.get("title") or "").strip() or path.stem
            author = (metadata.get("author") or "").strip() or None
            page_count = document.page_count
        finally:
            document.close()

        if not any(page.strip() for page in pages):
            raise IngestionError(
                f"No extractable text in {path.name} — it is probably a scanned image PDF.",
                hint="Run OCR on the file first, then import the OCR'd version.",
            )

        return LoadedDocument(
            text="\n\n".join(pages),
            title=title,
            author=author,
            doc_type=self.doc_type,
            meta={
                "pages": page_count,
                "empty_pages": empty_pages,
                "producer": (metadata.get("producer") or "").strip() or None,
            },
            sections=pages,
        )


class DocxLoader(DocumentLoader):
    extensions = {".docx"}
    doc_type = "docx"

    def load(self, path: Path, *, max_rows: int | None = None) -> LoadedDocument:
        try:
            import docx
        except ImportError as exc:
            raise DependencyMissingError("python-docx", "Word import", install="python-docx") from exc

        try:
            document = docx.Document(str(path))
        except Exception as exc:  # noqa: BLE001
            raise IngestionError(
                f"Could not open {path.name}: {exc}",
                hint="Only .docx is supported — re-save legacy .doc files as .docx.",
            ) from exc

        blocks: list[str] = []
        headings = 0
        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if not text:
                continue
            style = (paragraph.style.name or "").lower() if paragraph.style else ""
            if style.startswith("heading"):
                headings += 1
                level = "".join(ch for ch in style if ch.isdigit()) or "1"
                blocks.append(f"{'#' * min(int(level), 6)} {text}")
            else:
                blocks.append(text)

        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    blocks.append(" | ".join(cells))

        properties = document.core_properties
        return LoadedDocument(
            text="\n\n".join(blocks),
            title=(properties.title or "").strip() or path.stem,
            author=(properties.author or "").strip() or None,
            doc_type=self.doc_type,
            meta={"paragraphs": len(blocks), "headings": headings, "tables": len(document.tables)},
            sections=blocks,
        )


class EPUBLoader(DocumentLoader):
    extensions = {".epub"}
    doc_type = "epub"

    def load(self, path: Path, *, max_rows: int | None = None) -> LoadedDocument:
        try:
            import ebooklib
            from ebooklib import epub
        except ImportError as exc:
            raise DependencyMissingError("ebooklib", "EPUB import", install="ebooklib") from exc

        from ai_studio.data.loaders.text_loaders import html_to_document

        try:
            book = epub.read_epub(str(path))
        except Exception as exc:  # noqa: BLE001
            raise IngestionError(f"Could not open {path.name}: {exc}") from exc

        chapters: list[str] = []
        for item in book.get_items():
            if item.get_type() != ebooklib.ITEM_DOCUMENT:
                continue
            try:
                raw = item.get_content().decode("utf-8", errors="replace")
                parsed = html_to_document(raw, doc_type="epub")
                if parsed.text.strip():
                    chapters.append(parsed.text.strip())
            except DependencyMissingError:
                raise
            except Exception:  # noqa: BLE001 - skip an unreadable chapter, keep the book
                continue

        if not chapters:
            raise IngestionError(f"No readable chapters found in {path.name}")

        def _meta(field: str) -> str | None:
            try:
                values = book.get_metadata("DC", field)
                return str(values[0][0]).strip() if values else None
            except Exception:  # noqa: BLE001
                return None

        return LoadedDocument(
            text="\n\n".join(chapters),
            title=_meta("title") or path.stem,
            author=_meta("creator"),
            doc_type=self.doc_type,
            meta={
                "chapters": len(chapters),
                "language": _meta("language"),
                "publisher": _meta("publisher"),
            },
            sections=chapters,
        )


def strip_soft_hyphens(text: str) -> str:
    return re.sub(r"­", "", text)
